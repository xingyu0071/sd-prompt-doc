# -*- coding: utf-8 -*-
"""
扫描指定文件夹中的图片，识别 stable-diffusion-webui-aki（A1111 系）写入的
提示词元数据，并把「图片 + 正向提示词 + 负向提示词 + 生成参数」汇总成一个
Word 文档（.docx）。

用法示例：
    python sd_aki_prompt_doc.py "D:\\画图输出"
    python sd_aki_prompt_doc.py "D:\\画图输出" -o "D:\\提示词图集.docx" --page-break
"""

from __future__ import annotations

import argparse
import os
import re
import sys
import tempfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

# ---------------------------------------------------------------- 依赖检查

try:
    from PIL import Image
except ImportError:  # pragma: no cover
    sys.exit("缺少依赖 Pillow，请先运行 setup_env.bat，或执行：pip install pillow")

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:  # pragma: no cover
    sys.exit("缺少依赖 python-docx，请先运行 setup_env.bat，或执行：pip install python-docx")


IMAGE_EXTS = {
    ".png", ".jpg", ".jpeg", ".jpe", ".jfif",
    ".webp", ".bmp", ".gif", ".tif", ".tiff",
}

STEPS_RE = re.compile(r"(^|\n)\s*Steps:\s*\d+", re.IGNORECASE)
NEG_PREFIX = "Negative prompt:"


# ---------------------------------------------------------------- 元数据读取

@dataclass
class ImageInfo:
    path: Path
    rel_path: str = ""
    has_prompt: bool = False
    is_aki: bool = False
    prompt: str = ""
    negative: str = ""
    settings: str = ""
    source: str = ""
    format_name: str = "无"
    note: str = ""
    error: str = ""
    extra: dict = field(default_factory=dict)


def _decode_user_comment(raw) -> str:
    """解码 EXIF UserComment（A1111 在 jpg/webp 里用它存参数）。"""
    if raw is None:
        return ""
    if isinstance(raw, str):
        return raw
    try:
        data = bytes(raw)
    except Exception:
        return ""
    if data.startswith(b"UNICODE\x00"):
        return data[8:].decode("utf-16-be", errors="ignore").rstrip("\x00")
    if data.startswith(b"ASCII\x00\x00\x00"):
        return data[8:].decode("ascii", errors="ignore").rstrip("\x00")
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="ignore").rstrip("\x00")
    for enc in ("utf-8", "utf-16-be", "gbk", "latin-1"):
        try:
            text = data.decode(enc).rstrip("\x00")
        except Exception:
            continue
        if "Steps:" in text or "Negative prompt:" in text:
            return text
    return data.decode("utf-8", errors="ignore").rstrip("\x00")


def _read_exif_user_comment(img: Image.Image) -> str:
    try:
        exif = img.getexif()
    except Exception:
        return ""
    if not exif:
        return ""
    value = exif.get(0x9286)  # UserComment
    if value is None:
        try:
            sub = exif.get_ifd(0x8769)
            value = sub.get(0x9286)
        except Exception:
            value = None
    return _decode_user_comment(value)


def parse_aki_parameters(text: str) -> tuple[str, str, str]:
    """拆分 A1111/Aki 的 parameters 文本 -> (正向, 负向, 参数行)。"""
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip("\n")
    lines = normalized.split("\n")

    settings_idx = None
    for i in range(len(lines) - 1, -1, -1):
        if re.match(r"^\s*Steps:\s*\d+", lines[i], re.IGNORECASE):
            settings_idx = i
            break

    if settings_idx is None:
        settings = ""
        body = lines
    else:
        settings = lines[settings_idx].strip()
        body = lines[:settings_idx]

    while body and not body[-1].strip():
        body.pop()

    neg_idx = None
    for i in range(len(body) - 1, -1, -1):
        if body[i].lstrip().startswith(NEG_PREFIX):
            neg_idx = i
            break

    if neg_idx is None:
        return "\n".join(body).strip(), "", settings

    positive = "\n".join(body[:neg_idx]).strip()
    first = body[neg_idx].lstrip()[len(NEG_PREFIX):].lstrip()
    negative = "\n".join([first] + body[neg_idx + 1:]).strip()
    return positive, negative, settings


def read_image_info(path: Path, root: Path) -> ImageInfo:
    info = ImageInfo(path=path)
    try:
        info.rel_path = str(path.relative_to(root))
    except ValueError:
        info.rel_path = path.name

    try:
        with Image.open(path) as img:
            meta = dict(img.info or {})
            exif_text = _read_exif_user_comment(img)
            comfy_prompt = meta.get("prompt") or meta.get("Prompt")
    except Exception as exc:  # 损坏 / 非图片 / 权限
        info.error = f"{type(exc).__name__}: {exc}"
        return info

    raw = ""
    source = ""
    for key in ("parameters", "Parameters"):
        if isinstance(meta.get(key), str) and meta[key].strip():
            raw, source = meta[key], f"PNG tEXt: {key}"
            break
    if not raw:
        for key in ("Comment", "comment", "Description", "description"):
            if isinstance(meta.get(key), str) and meta[key].strip():
                raw, source = meta[key], f"PNG tEXt: {key}"
                break
    if not raw and exif_text.strip():
        raw, source = exif_text, "EXIF UserComment"

    # ComfyUI 的 prompt 是 JSON 工作流，不是 Aki 格式
    if not raw and isinstance(comfy_prompt, str) and comfy_prompt.strip().startswith("{"):
        info.format_name = "ComfyUI 工作流"
        info.note = "含 ComfyUI workflow，非 Aki 提示词"
        info.extra["comfy"] = True
        return info

    if not raw:
        info.format_name = "无"
        info.note = "图片中没有提示词元数据"
        return info

    info.source = source
    positive, negative, settings = parse_aki_parameters(raw)

    if STEPS_RE.search(raw):
        info.is_aki = True
        info.format_name = "Aki / A1111"
    else:
        info.format_name = "其他格式"
        info.note = "缺少 Steps 参数行，可能不是 Aki 生成的"

    info.prompt = positive
    info.negative = negative
    info.settings = settings
    info.has_prompt = bool(positive.strip())
    if not info.has_prompt:
        info.note = info.note or "元数据里没有可用的正向提示词"
    return info


# ---------------------------------------------------------------- Word 排版

def shade_paragraph(paragraph, fill: str = "F3F4F6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 9) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    _set_east_asia(normal, "微软雅黑")

    for style_name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        style = doc.styles[style_name]
        style.font.name = "微软雅黑"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0x1F, 0x3B, 0x57)
        _set_east_asia(style, "微软雅黑")


def _set_east_asia(style, font_name: str) -> None:
    """同时设置西文与中文字体，避免中文回退成宋体。"""
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    # 主题字体优先级高于显式字体，必须去掉，否则中文会回退成默认主题字体
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if r_fonts.get(qn(attr)) is not None:
            del r_fonts.attrib[qn(attr)]


def add_label(doc: Document, text: str, color: str = "1F4E79") -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(9.5)
    run.font.name = "微软雅黑"
    run.font.color.rgb = RGBColor.from_string(color)


def add_block_text(doc: Document, text: str, fill: str = "F3F4F6") -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.2)
    shade_paragraph(p, fill)
    run = p.add_run(text if text else "（空）")
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if not text:
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)


def prepare_picture(path: Path, max_side: int, tmpdir: Path):
    """返回可交给 python-docx 的图片路径；必要时转格式/缩放。"""
    with Image.open(path) as img:
        if getattr(img, "is_animated", False):
            img.seek(0)
        width, height = img.size
        biggest = max(width, height)
        need_resize = bool(max_side and biggest > max_side)
        fmt = (img.format or "").upper()
        has_alpha = img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info)
        supported = fmt in ("PNG", "JPEG", "GIF", "BMP", "TIFF")

        if supported and not need_resize:
            return path, False

        img.load()
        if need_resize:
            ratio = max_side / float(biggest)
            new_size = (max(1, int(width * ratio)), max(1, int(height * ratio)))
            img = img.resize(new_size, Image.LANCZOS)

        if has_alpha or fmt == "PNG":
            out = tmpdir / f"{abs(hash((str(path), os.getpid())))}.png"
            if img.mode not in ("RGB", "RGBA"):
                img = img.convert("RGBA" if has_alpha else "RGB")
            img.save(out, "PNG", optimize=True)
        else:
            out = tmpdir / f"{abs(hash((str(path), os.getpid())))}.jpg"
            img.convert("RGB").save(out, "JPEG", quality=92, optimize=True)
        return out, True


def build_document(infos: list[ImageInfo], output: Path, args) -> Path:
    doc = Document()
    setup_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("SD WebUI（Aki）提示词图集")
    run.bold = True
    run.font.size = Pt(22)
    run.font.name = "微软雅黑"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x57)

    total = len(infos)
    with_prompt = sum(1 for i in infos if i.has_prompt)
    aki_count = sum(1 for i in infos if i.is_aki)
    failed = sum(1 for i in infos if i.error)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        f"图片文件夹：{args.folder}\n"
        f"扫描图片 {total} 张 ｜ 含提示词 {with_prompt} 张 ｜ Aki 格式 {aki_count} 张"
        f" ｜ 无法读取 {failed} 张\n生成时间：{datetime.now():%Y-%m-%d %H:%M:%S}"
    )
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    if args.summary_table and infos:
        doc.add_paragraph()
        add_label(doc, "扫描结果一览")
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ("#", "文件名", "是否含提示词", "格式 / 备注")
        for cell, text in zip(table.rows[0].cells, headers):
            set_cell_text(cell, text, bold=True)
            shade_paragraph(cell.paragraphs[0], "DCE6F1")
        for idx, item in enumerate(infos, 1):
            if item.error:
                status = "读取失败"
            elif item.has_prompt:
                status = "是"
            else:
                status = "否"
            remark = item.format_name if not item.note else f"{item.format_name}；{item.note}"
            if item.error:
                remark = item.error
            cells = table.add_row().cells
            set_cell_text(cells[0], str(idx))
            set_cell_text(cells[1], item.rel_path)
            set_cell_text(cells[2], status)
            set_cell_text(cells[3], remark)
        widths = (Cm(1.0), Cm(7.2), Cm(2.6), Cm(6.2))
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = width

    doc.add_paragraph()
    doc.add_paragraph("提示词正文", style="Heading 2")

    index = 0
    with tempfile.TemporaryDirectory(prefix="sd_prompt_doc_") as tmp:
        tmpdir = Path(tmp)
        for item in infos:
            if not item.has_prompt:
                continue
            index += 1
            if args.page_break and index > 1:
                doc.add_page_break()

            heading = doc.add_paragraph()
            heading.paragraph_format.space_before = Pt(10)
            head_run = heading.add_run(f"{index}. {item.rel_path}")
            head_run.bold = True
            head_run.font.size = Pt(13)
            head_run.font.name = "微软雅黑"
            head_run.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

            info_line = doc.add_paragraph()
            info_run = info_line.add_run(
                f"格式：{item.format_name}　｜　来源：{item.source or '未知'}"
            )
            info_run.font.size = Pt(8)
            info_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

            try:
                pic_path, _ = prepare_picture(item.path, args.max_side, tmpdir)
                pic_para = doc.add_paragraph()
                pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic_run = pic_para.add_run()
                pic_run.add_picture(str(pic_path), width=Cm(args.max_width_cm))
            except Exception as exc:
                warn = doc.add_paragraph()
                warn_run = warn.add_run(f"[图片插入失败] {type(exc).__name__}: {exc}")
                warn_run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
                warn_run.font.size = Pt(9)

            add_label(doc, "正向提示词")
            add_block_text(doc, item.prompt, "EEF4FB")

            if item.negative:
                add_label(doc, "负向提示词")
                add_block_text(doc, item.negative, "FBEEEE")

            if item.settings:
                add_label(doc, "生成参数")
                add_block_text(doc, item.settings, "F3F4F6")

    no_prompt = [i for i in infos if not i.has_prompt and not i.error]
    failed_items = [i for i in infos if i.error]
    if no_prompt or failed_items:
        doc.add_page_break()
        doc.add_paragraph("未收录进正文的图片", style="Heading 2")

        if no_prompt:
            note = doc.add_paragraph()
            note_run = note.add_run(
                "以下图片没有检测到 Aki/A1111 提示词元数据（可能被压缩、转存或截图时丢失了信息）："
            )
            note_run.font.size = Pt(9)
            note_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            for item in no_prompt:
                p = doc.add_paragraph(style="List Bullet")
                text = item.rel_path
                if item.format_name != "无":
                    text += f"（{item.format_name}）"
                p.add_run(text).font.size = Pt(9)

        if failed_items:
            note = doc.add_paragraph()
            note_run = note.add_run("以下文件无法作为图片读取：")
            note_run.font.size = Pt(9)
            note_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            for item in failed_items:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f"{item.rel_path}：{item.error}").font.size = Pt(9)

    core = doc.core_properties
    core.title = "SD WebUI（Aki）提示词图集"
    core.author = "sd_aki_prompt_doc.py"
    core.comments = f"由 {args.folder} 扫描生成"

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return output


# ---------------------------------------------------------------- 主流程

def collect_images(folder: Path, recursive: bool, only_aki: bool) -> list[ImageInfo]:
    iterator = folder.rglob("*") if recursive else folder.glob("*")
    files = sorted(
        (p for p in iterator if p.is_file() and p.suffix.lower() in IMAGE_EXTS),
        key=lambda p: (p.parent.as_posix().lower(), p.name.lower()),
    )
    infos = [read_image_info(path, folder) for path in files]
    if only_aki:
        infos = [i for i in infos if i.is_aki]
    return infos


def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description="扫描图片中的 sd-webui-aki 提示词，并生成图片+提示词 Word 文档。",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("folder", help="存放图片的文件夹")
    parser.add_argument("-o", "--output", default=None, help="输出 docx 路径")
    parser.add_argument("--no-recursive", action="store_true", help="只扫描当前文件夹，不含子文件夹")
    parser.add_argument("--only-aki", action="store_true", help="只保留 Aki/A1111 格式的图片")
    parser.add_argument("--max-width-cm", type=float, default=15.0, help="文档中图片宽度（厘米）")
    parser.add_argument("--max-side", type=int, default=1600,
                        help="嵌入前把图片长边缩到此像素，0 表示不缩放")
    parser.add_argument("--page-break", action="store_true", help="每张图片单独起一页")
    parser.add_argument("--no-summary-table", dest="summary_table", action="store_false",
                        help="不生成扫描结果一览表")
    parser.add_argument("--list-only", action="store_true", help="只打印识别结果，不生成文档")
    parser.set_defaults(summary_table=True)
    return parser.parse_args(argv)


def main(argv=None) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    args = parse_args(argv)
    # 容忍拖拽/手输时多带的引号和首尾空格
    raw_folder = args.folder.strip().strip('"').strip()
    folder = Path(raw_folder).expanduser()
    if not folder.is_dir():
        print(f"[错误] 文件夹不存在：{folder}")
        return 2

    args.folder = str(folder.resolve())
    print(f"[1/3] 扫描文件夹：{folder}")
    infos = collect_images(folder, not args.no_recursive, args.only_aki)

    if not infos:
        print("[提示] 没有找到图片（支持 png/jpg/jpeg/webp/bmp/gif/tif）。")
        return 1

    with_prompt = sum(1 for i in infos if i.has_prompt)
    aki = sum(1 for i in infos if i.is_aki)
    failed = [i for i in infos if i.error]
    print(f"      共 {len(infos)} 张图片，含提示词 {with_prompt} 张，Aki 格式 {aki} 张，读取失败 {len(failed)} 张")

    if args.list_only:
        for item in infos:
            flag = "AKI " if item.is_aki else ("PROMPT" if item.has_prompt else "  -   ")
            print(f"  [{flag}] {item.rel_path}" + (f"  <- {item.note or item.error}" if (item.note or item.error) else ""))
        return 0

    if args.output:
        output = Path(args.output).expanduser()
        if output.suffix.lower() != ".docx":
            output = output.with_suffix(".docx")
    else:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output = folder / f"SD提示词图集_{stamp}.docx"

    print(f"[2/3] 生成文档：{output}")
    build_document(infos, output, args)

    size_mb = output.stat().st_size / 1024 / 1024
    print(f"[3/3] 完成：{output}（{size_mb:.2f} MB）")
    if failed:
        print("      以下图片读取失败：")
        for item in failed:
            print(f"        - {item.rel_path}: {item.error}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
